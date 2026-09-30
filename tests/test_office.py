import base64
from io import BytesIO
from pathlib import Path

import pytest

from aura.docedit import handle_doc_image_edit, wants_doc_image_edit
from aura.readers import ReadResult, read_bytes


def _minimal_docx(text: str = "Olá mundo") -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _minimal_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"
    ws["A1"] = "nome"
    ws["B1"] = "idade"
    ws["A2"] = "Ana"
    ws["B2"] = 30
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _minimal_pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Título"
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_read_xlsx_and_pptx():
    x = read_bytes("planilha.xlsx", _minimal_xlsx())
    assert x.kind == "document"
    assert "Ana" in x.text
    p = read_bytes("slides.pptx", _minimal_pptx())
    assert p.kind == "document"
    assert "Título" in p.text


def test_pdf_with_images():
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Sensor A1")
    # tiny inline image
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 8, 8), 1)
    page.insert_image(pymupdf.Rect(72, 100, 120, 148), pixmap=pix)
    data = doc.tobytes()
    doc.close()
    result = read_bytes("circuitos.pdf", data)
    assert result.kind in ("pdf_mixed", "pdf_scan")
    assert result.pages_b64


def test_read_pdf_text():
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Relatorio mensal de vendas - janeiro 2026")
    data = doc.tobytes()
    doc.close()
    result = read_bytes("relatorio.pdf", data)
    assert result.kind == "document"
    assert "vendas" in result.text


def test_pdf_scan_flag(monkeypatch):
    def fake_pages(_data: bytes, max_pages: int = 10) -> list[str]:
        return ["ZmFrZQ=="]

    monkeypatch.setattr("aura.readers._pdf_pages_as_b64", fake_pages)
    import pymupdf

    doc = pymupdf.open()
    doc.new_page()
    data = doc.tobytes()
    doc.close()
    result = read_bytes("scan.pdf", data)
    assert result.kind == "pdf_scan"
    assert result.pages_b64


def test_wants_doc_image_edit():
    assert wants_doc_image_edit("insira foto.png no relatorio.docx")
    assert wants_doc_image_edit("remova as imagens do documento nota.docx")
    assert not wants_doc_image_edit("leia relatorio.docx")


def test_insert_and_remove_docx_images(tmp_path, monkeypatch):
    doc = tmp_path / "relatorio.docx"
    doc.write_bytes(_minimal_docx())
    img = tmp_path / "logo.png"
    img.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    ))
    ok, msg = handle_doc_image_edit(f"insira {img} no {doc}")
    assert ok, msg
    from docx import Document

    loaded = Document(doc)
    assert len(loaded.inline_shapes) >= 1
    ok2, msg2 = handle_doc_image_edit(f"remova imagens de {doc}")
    assert ok2, msg2
    loaded2 = Document(doc)
    assert len(loaded2.inline_shapes) == 0
