from __future__ import annotations

import io
import sys
import tempfile
from pathlib import Path

_MAX = 24_000


def _clip(text: str) -> str:
    text = text.strip()
    if len(text) > _MAX:
        return text[:_MAX] + "\n\n[conteúdo truncado]"
    return text


def com_available() -> bool:
    return sys.platform == "win32"


def read_xlsx_bytes(data: bytes) -> tuple[str, str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets[:6]:
        lines.append(f"## {sheet.title}")
        for row in sheet.iter_rows(max_row=250, values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if any(cells):
                lines.append("\t".join(cells))
    wb.close()
    return _clip("\n".join(lines)), "Excel"


def read_pptx_bytes(data: bytes) -> tuple[str, str]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    total = len(prs.slides)
    for i in range(min(total, 40)):
        slide = prs.slides[i]
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            parts.append(f"--- Slide {i + 1} ---\n" + "\n".join(texts))
    return _clip("\n\n".join(parts)), "PowerPoint"


def _com_text(path: Path, app: str, open_fmt: int | None = None) -> str:
    import win32com.client  # type: ignore[import-untyped]

    prog = win32com.client.Dispatch(app)
    prog.Visible = False
    try:
        if open_fmt is not None:
            doc = prog.Documents.Open(str(path.resolve()), ReadOnly=True, Format=open_fmt)
        else:
            doc = prog.Documents.Open(str(path.resolve()), ReadOnly=True)
        if app.endswith("Word.Application"):
            text = str(doc.Content.Text or "")
        elif app.endswith("Excel.Application"):
            used = doc.UsedRange
            rows = int(used.Rows.Count)
            cols = int(used.Columns.Count)
            lines: list[str] = []
            for r in range(1, min(rows, 250) + 1):
                row_vals = []
                for c in range(1, min(cols, 30) + 1):
                    v = used.Cells(r, c).Value
                    row_vals.append("" if v is None else str(v).strip())
                if any(row_vals):
                    lines.append("\t".join(row_vals))
            text = "\n".join(lines)
            doc.Close(False)
            return text.strip()
        else:
            parts: list[str] = []
            for i in range(1, min(int(doc.Slides.Count), 40) + 1):
                slide = doc.Slides(i)
                texts: list[str] = []
                for j in range(1, int(slide.Shapes.Count) + 1):
                    shp = slide.Shapes(j)
                    if shp.HasTextFrame:
                        t = shp.TextFrame.TextRange.Text
                        if t and str(t).strip():
                            texts.append(str(t).strip())
                if texts:
                    parts.append(f"--- Slide {i} ---\n" + "\n".join(texts))
            text = "\n\n".join(parts)
        doc.Close(False)
        return text.strip()
    finally:
        prog.Quit()


def read_legacy_path(path: Path) -> tuple[str, str]:
    ext = path.suffix.lower()
    if not com_available():
        raise RuntimeError(f"{ext} precisa do Microsoft Office no Windows.")
    if ext == ".doc":
        return _clip(_com_text(path, "Word.Application", open_fmt=0)), "Word (.doc)"
    if ext == ".xls":
        return _clip(_com_text(path, "Excel.Application")), "Excel (.xls)"
    if ext == ".ppt":
        return _clip(_com_text(path, "PowerPoint.Application")), "PowerPoint (.ppt)"
    raise RuntimeError(f"Formato {ext} legado não suportado.")


def read_legacy_bytes(name: str, data: bytes) -> tuple[str, str]:
    ext = Path(name).suffix.lower()
    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        tmp.write(data)
        tmp_path = Path(tmp.name)
    try:
        return read_legacy_path(tmp_path)
    finally:
        try:
            tmp_path.unlink(missing_ok=True)
        except OSError:
            pass
