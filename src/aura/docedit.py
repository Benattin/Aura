from __future__ import annotations

import base64
import re
import tempfile
from pathlib import Path

from aura.office import com_available
from aura.skills import resolve_existing, search_files

_INSERT = re.compile(
    r"\b(insira|insere|inserir|adicion\w*|coloc\w*|ponh\w*|inclu\w*)\b.*"
    r"(?:\b(imagem|foto|figura)\b|[\w.\-]+\.(?:png|jpe?g|gif|webp|bmp))",
    re.I,
)
_REMOVE = re.compile(
    r"\b(apag\w*|remov\w*|exclu\w*|tir\w*|delete)\b.*\b(imagem|foto|figura|imagens|fotos|figuras)\b",
    re.I,
)
_DOC = re.compile(r"([^\s\"']+\.(?:docx|doc))", re.I)
_IMG = re.compile(
    r"([a-zA-Z]:[\\/][^\s\"']+\.(?:png|jpe?g|gif|webp|bmp)"
    r"|[^\s\"']+\.(?:png|jpe?g|gif|webp|bmp))",
    re.I,
)


def wants_doc_image_edit(text: str) -> bool:
    t = text.strip()
    if not (_INSERT.search(t) or _REMOVE.search(t)):
        return False
    return bool(_DOC.search(t) or re.search(r"\b(documento|word|docx?)\b", t, re.I))


def _resolve_file(query: str) -> Path | None:
    if not query:
        return None
    direct = resolve_existing(query)
    if direct and direct.is_file():
        return direct
    hits = [p for p in search_files(query) if p.is_file()]
    return hits[0] if hits else None


def _doc_path(text: str) -> Path | None:
    if m := _DOC.search(text):
        p = _resolve_file(m.group(1))
        if p:
            return p
        candidate = Path(m.group(1))
        if candidate.is_file():
            return candidate
    for token in re.findall(r"[\w.\-\\/:\s]+\.(?:docx|doc)", text, re.I):
        token = token.strip()
        if p := _resolve_file(token):
            return p
        candidate = Path(token)
        if candidate.is_file():
            return candidate
    if m := re.search(r"\b(documento|word)\s+([\w.\-]+)", text, re.I):
        return _resolve_file(m.group(2))
    return None


def _image_path(text: str) -> Path | None:
    if m := _IMG.search(text):
        p = _resolve_file(m.group(1))
        if p:
            return p
        candidate = Path(m.group(1))
        if candidate.is_file():
            return candidate
    return None


def wants_insert(text: str) -> bool:
    return bool(_INSERT.search(text))


def wants_remove_images(text: str) -> bool:
    return bool(_REMOVE.search(text))


def _insert_docx(doc: Path, image: Path) -> tuple[bool, str]:
    from docx import Document
    from docx.shared import Inches

    document = Document(doc)
    document.add_picture(str(image), width=Inches(5.5))
    document.save(doc)
    return True, f"Inseri {image.name} em {doc}, senhor."


def _insert_doc_com(doc: Path, image: Path) -> tuple[bool, str]:
    import win32com.client  # type: ignore[import-untyped]

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        opened = word.Documents.Open(str(doc.resolve()))
        opened.InlineShapes.AddPicture(FileName=str(image.resolve()))
        opened.Save()
        opened.Close()
        return True, f"Inseri {image.name} em {doc}, senhor."
    finally:
        word.Quit()


def _remove_docx_images(doc: Path) -> tuple[bool, str]:
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(doc)
    count = 0
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            drawings = run._element.findall(".//" + qn("w:drawing"))
            count += len(drawings)
            for drawing in drawings:
                parent = drawing.getparent()
                if parent is not None:
                    parent.remove(drawing)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        drawings = run._element.findall(".//" + qn("w:drawing"))
                        count += len(drawings)
                        for drawing in drawings:
                            parent = drawing.getparent()
                            if parent is not None:
                                parent.remove(drawing)
    document.save(doc)
    if count == 0:
        return True, f"Não havia imagens em {doc.name}, senhor."
    return True, f"Removi {count} imagem(ns) de {doc.name}, senhor."


def _remove_doc_com(doc: Path) -> tuple[bool, str]:
    import win32com.client  # type: ignore[import-untyped]

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    count = 0
    try:
        opened = word.Documents.Open(str(doc.resolve()))
        while opened.InlineShapes.Count > 0:
            opened.InlineShapes(1).Delete()
            count += 1
        opened.Save()
        opened.Close()
        if count == 0:
            return True, f"Não havia imagens em {doc.name}, senhor."
        return True, f"Removi {count} imagem(ns) de {doc.name}, senhor."
    finally:
        word.Quit()


def insert_image(doc: Path, image: Path) -> tuple[bool, str]:
    if not doc.exists():
        return False, f"Não encontrei o documento {doc.name}."
    if not image.exists():
        return False, f"Não encontrei a imagem {image.name}."
    ext = doc.suffix.lower()
    try:
        if ext == ".docx":
            return _insert_docx(doc, image)
        if ext == ".doc" and com_available():
            return _insert_doc_com(doc, image)
        return False, "Só edito imagens em .docx ou .doc com Word instalado."
    except Exception as exc:
        return False, f"Não consegui inserir a imagem: {exc}"


def remove_images(doc: Path) -> tuple[bool, str]:
    if not doc.exists():
        return False, f"Não encontrei o documento {doc.name}."
    ext = doc.suffix.lower()
    try:
        if ext == ".docx":
            return _remove_docx_images(doc)
        if ext == ".doc" and com_available():
            return _remove_doc_com(doc)
        return False, "Só removo imagens de .docx ou .doc com Word instalado."
    except Exception as exc:
        return False, f"Não consegui remover imagens: {exc}"


def insert_image_bytes(doc_text: str, image_name: str, raw: bytes) -> tuple[bool, str]:
    doc = _doc_path(doc_text)
    if not doc:
        return False, "Qual documento Word devo editar? Diga o nome ou caminho do .docx."
    suffix = Path(image_name).suffix or ".png"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(raw)
        img_path = Path(tmp.name)
    try:
        return insert_image(doc, img_path)
    finally:
        try:
            img_path.unlink(missing_ok=True)
        except OSError:
            pass


def handle_doc_image_edit(text: str) -> tuple[bool, str]:
    doc = _doc_path(text)
    if not doc:
        return False, "Qual documento Word? Diga algo como: insira foto.png no relatorio.docx"
    if wants_remove_images(text):
        return remove_images(doc)
    image = _image_path(text)
    if not image:
        return False, "Qual imagem devo inserir? Diga o nome ou caminho do arquivo de imagem."
    return insert_image(doc, image)


def insert_from_b64(text: str, image_name: str, data_b64: str) -> tuple[bool, str]:
    try:
        raw = base64.b64decode(data_b64, validate=False)
    except Exception:
        return False, "Imagem anexada corrompida, senhor."
    return insert_image_bytes(text, image_name, raw)
